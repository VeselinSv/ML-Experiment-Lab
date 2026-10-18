
"""Зареждане и първоначална проверка на набори от данни."""

from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path
from zipfile import BadZipFile


class DatasetLoadError(ValueError):
    """Грешка при зареждане или проверка на набор от данни."""


@dataclass(frozen=True)
class LoadedDataset:
    """Набор от данни, зареден в паметта."""

    path: Path
    columns: tuple[str, ...]
    rows: tuple[tuple[str, ...], ...]
    sheet_name: str | None = None

    @property
    def row_count(self) -> int:
        """Брой редове с данни, без заглавния ред."""
        return len(self.rows)

    @property
    def column_count(self) -> int:
        """Брой колони."""
        return len(self.columns)

    def preview(
        self,
        limit: int = 5,
    ) -> tuple[tuple[str, ...], ...]:
        """Връща до limit реда за предварителен преглед."""
        if limit < 0:
            raise ValueError("limit must be zero or greater")

        return self.rows[:limit]


def _validate_file_path(
    file_path: str | Path,
) -> Path:
    """Проверява дали пътят сочи към съществуващ файл."""
    path = Path(file_path).expanduser()

    if not path.exists():
        raise DatasetLoadError(f"File does not exist: {path}")

    if not path.is_file():
        raise DatasetLoadError(
            f"The selected path is not a file: {path}"
        )

    return path


def _validate_headers(
    raw_headers: list[str],
) -> tuple[str, ...]:
    """Проверява заглавния ред и връща имената на колоните."""
    columns = tuple(header.strip() for header in raw_headers)

    if not columns or all(not column for column in columns):
        raise DatasetLoadError(
            "The first row must contain column names."
        )

    if any(not column for column in columns):
        raise DatasetLoadError(
            "Every column must have a name."
        )

    if len(set(columns)) != len(columns):
        raise DatasetLoadError(
            "Column names must be unique."
        )

    return columns


def load_csv(file_path: str | Path) -> LoadedDataset:
    """Зарежда CSV файл и проверява структурата му."""
    path = _validate_file_path(file_path)

    if path.suffix.casefold() != ".csv":
        raise DatasetLoadError(
            "The selected file must have a .csv extension."
        )

    try:
        with path.open(
            "r",
            encoding="utf-8-sig",
            newline="",
        ) as file:
            sample = file.read(4096)

            if not sample.strip():
                raise DatasetLoadError(
                    "The CSV file is empty."
                )

            file.seek(0)

            try:
                dialect = csv.Sniffer().sniff(
                    sample,
                    delimiters=",;\t|",
                )
            except csv.Error:
                dialect = csv.excel

            reader = csv.reader(file, dialect)
            raw_headers = next(reader, None)

            if raw_headers is None:
                raise DatasetLoadError(
                    "The CSV file is empty."
                )

            columns = _validate_headers(raw_headers)
            rows: list[tuple[str, ...]] = []

            for record_number, raw_row in enumerate(
                reader,
                start=2,
            ):
                # Празните редове не се добавят към набора.
                if not raw_row or all(
                    not value.strip() for value in raw_row
                ):
                    continue

                if len(raw_row) != len(columns):
                    raise DatasetLoadError(
                        f"Invalid record {record_number}: "
                        f"expected {len(columns)} values, "
                        f"but found {len(raw_row)}."
                    )

                rows.append(tuple(raw_row))

    except UnicodeDecodeError as exc:
        raise DatasetLoadError(
            "The file could not be read as UTF-8 text. "
            "Save it as UTF-8 and try again."
        ) from exc

    except csv.Error as exc:
        raise DatasetLoadError(
            f"The CSV format is invalid: {exc}"
        ) from exc

    except OSError as exc:
        raise DatasetLoadError(
            f"Could not open the CSV file: {exc}"
        ) from exc

    return LoadedDataset(
        path=path.resolve(),
        columns=columns,
        rows=tuple(rows),
    )


def get_excel_sheet_names(
    file_path: str | Path,
) -> tuple[str, ...]:
    """Връща имената на работните листове в Excel файл."""
    path = _validate_file_path(file_path)

    if path.suffix.casefold() != ".xlsx":
        raise DatasetLoadError(
            "The selected file must have an .xlsx extension."
        )

    try:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError as exc:
        raise DatasetLoadError(
            "Excel import requires openpyxl. "
            "Install the project dependencies and try again."
        ) from exc

    try:
        workbook = load_workbook(
            path,
            read_only=True,
            data_only=True,
        )
    except (OSError, BadZipFile, InvalidFileException, ValueError) as exc:
        raise DatasetLoadError(
            f"Could not open the Excel workbook: {exc}"
        ) from exc

    try:
        return tuple(workbook.sheetnames)
    finally:
        workbook.close()


def _excel_value_to_text(value: object) -> str:
    """Преобразува стойност от Excel клетка в текст."""
    if value is None:
        return ""

    return str(value)


def load_xlsx(
    file_path: str | Path,
    sheet_name: str | None = None,
) -> LoadedDataset:
    """Зарежда избран работен лист от Excel файл.

    Ако sheet_name не е зададен, се използва активният работен лист.
    Празните редове се пропускат, а липсващите клетки се представят
    като празни текстови стойности.
    """
    path = _validate_file_path(file_path)

    if path.suffix.casefold() != ".xlsx":
        raise DatasetLoadError(
            "The selected file must have an .xlsx extension."
        )

    try:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError as exc:
        raise DatasetLoadError(
            "Excel import requires openpyxl. "
            "Install the project dependencies and try again."
        ) from exc

    try:
        workbook = load_workbook(
            path,
            read_only=True,
            data_only=True,
        )
    except (OSError, BadZipFile, InvalidFileException, ValueError) as exc:
        raise DatasetLoadError(
            f"Could not open the Excel workbook: {exc}"
        ) from exc

    try:
        if not workbook.sheetnames:
            raise DatasetLoadError(
                "The Excel workbook does not contain worksheets."
            )

        if sheet_name is not None:
            if sheet_name not in workbook.sheetnames:
                raise DatasetLoadError(
                    f"Worksheet not found: {sheet_name}"
                )
            worksheet = workbook[sheet_name]
        else:
            worksheet = workbook.active

        if worksheet is None:
            raise DatasetLoadError(
                "The Excel workbook has no active worksheet."
            )

        max_column = worksheet.max_column or 1

        first_row = next(
            worksheet.iter_rows(
                min_row=1,
                max_row=1,
                max_col=max_column,
                values_only=True,
            ),
            None,
        )

        if first_row is None:
            raise DatasetLoadError(
                "The selected worksheet is empty."
            )

        raw_headers = list(first_row)

        # Премахваме празните клетки след последното име на колона.
        while raw_headers and (
            raw_headers[-1] is None
            or not str(raw_headers[-1]).strip()
        ):
            raw_headers.pop()

        if not raw_headers:
            raise DatasetLoadError(
                "The first row must contain column names."
            )

        columns = _validate_headers(
            [
                _excel_value_to_text(value)
                for value in raw_headers
            ]
        )

        rows: list[tuple[str, ...]] = []

        for record_number, raw_row in enumerate(
            worksheet.iter_rows(
                min_row=2,
                max_col=max_column,
                values_only=True,
            ),
            start=2,
        ):
            raw_values = list(raw_row)

            # Ако след заглавния ред има попълнена клетка без
            # съответна колона, файлът има структурно несъответствие.
            extra_values = raw_values[len(columns):]

            if any(
                value is not None and str(value).strip()
                for value in extra_values
            ):
                raise DatasetLoadError(
                    f"Invalid record {record_number}: "
                    "data exists beyond the named columns."
                )

            values = raw_values[:len(columns)]

            if len(values) < len(columns):
                values.extend(
                    [None] * (len(columns) - len(values))
                )

            text_row = tuple(
                _excel_value_to_text(value)
                for value in values
            )

            if all(not value.strip() for value in text_row):
                continue

            rows.append(text_row)

        return LoadedDataset(
            path=path.resolve(),
            columns=columns,
            rows=tuple(rows),
            sheet_name=worksheet.title,
        )

    finally:
        workbook.close()


def load_dataset(
    file_path: str | Path,
    sheet_name: str | None = None,
) -> LoadedDataset:
    """Зарежда поддържан набор от данни според файловото разширение."""
    path = _validate_file_path(file_path)
    suffix = path.suffix.casefold()

    if suffix == ".csv":
        return load_csv(path)

    if suffix == ".xlsx":
        return load_xlsx(path, sheet_name=sheet_name)

    raise DatasetLoadError(
        "Supported file types are .csv and .xlsx."
    )
